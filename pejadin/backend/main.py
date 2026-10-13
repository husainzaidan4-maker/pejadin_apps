import json,hmac,hashlib,base64,time,os,io,datetime as dt,logging
import tempfile
import psycopg
from psycopg.rows import dict_row
from psycopg_pool import ConnectionPool
from dotenv import load_dotenv
from fastapi import FastAPI,Depends,HTTPException,Request
from fastapi.responses import HTMLResponse,StreamingResponse,Response
from fastapi.staticfiles import StaticFiles
from pejadin_docs.business_rules import (
    Pegawai as PegawaiHitung,
    SPDHasil,
    SPDMasukan,
    SuratTugas as SuratTugasHitung,
    TarifSBM,
    UangRepresentasiRate,
    ValidationError as ErroValidasiHitung,
    hitung_spd,
)
from pejadin_docs.pdf_merge import gabungkan_pdf
from pejadin_docs.pdf_st import buat_pdf_st,buat_pdf_st_lampiran
from pejadin_docs.pdf_spd import buat_pdf_spd
from pejadin_docs.pdf_spd_hal2 import buat_pdf_spd_hal2
from pejadin_docs.pdf_rincian_kuitansi import buat_pdf_rincian,buat_pdf_kuitansi
from pejadin_docs.pdf_sptjb import buat_pdf_sptjb
from pejadin_docs.pdf_nominatif import buat_pdf_nominatif
from pejadin_docs.pdf_laporan import buat_pdf_laporan
from pejadin_docs.pdf_riil import buat_pdf_riil
from pejadin_docs.security import filter_spd_for_user, validate_token_secret
from pejadin_docs.satker_migration import (
    ensure_spd_pegawai_foreign_key,
    migrate_legacy_satker,
    sync_legacy_satker,
)
from pejadin_docs.rate_limit import RateLimiter, request_rate_limit
D=os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(D,".env"))
URL=os.getenv("DATABASE_URL","").strip()
SECRET=validate_token_secret(os.getenv("SECRET"))
app=FastAPI(title="PEJADIN")
logger=logging.getLogger(__name__)
rate_limiter=RateLimiter()
RATE_LIMIT_WINDOW_SECONDS=60
if not URL:raise RuntimeError("DATABASE_URL belum diisi. Salin backend/.env.example menjadi backend/.env lalu isi dengan connection string Supabase.")
pool=ConnectionPool(URL if "sslmode=" in URL else URL+("&" if "?" in URL else "?")+"sslmode=require",min_size=1,max_size=5,kwargs={"row_factory":dict_row,"prepare_threshold":None},open=False)

@app.middleware("http")
async def limit_sensitive_requests(request:Request,call_next):
    client=request.client.host if request.client else "unknown"
    path=request.url.path
    policy=request_rate_limit(request.method,path)
    if policy:
        scope,limit=policy
        if not rate_limiter.allow(
            f"{client}:{scope}",limit,RATE_LIMIT_WINDOW_SECONDS
        ):
            return Response(
                content=json.dumps({
                    "error":{
                        "code":"RATE_LIMITED",
                        "message":"Terlalu banyak permintaan. Coba lagi setelah beberapa saat.",
                    }
                }),
                status_code=429,
                media_type="application/json",
            )
    return await call_next(request)

def P(s):return s.replace("?","%s")
def q(sql,a=(),one=False):
    with pool.connection() as c:r=c.execute(P(sql),a).fetchall()
    return (r[0] if r else None) if one else r
def ex(sql,a=(),ret=False):
    with pool.connection() as c:
        cur=c.execute(P(sql)+(" returning id" if ret else ""),a)
        return cur.fetchone()["id"] if ret else None
INT="tahun harian inap1 inap2 inap3 inap4 taksi uang_harian_dalam_kota_8jam uang_harian_diklat urutan pegawai_id aktif".split()
def ci(b):
    try:return {k:(int(v) if k in INT and v not in(None,"") else v) for k,v in b.items()}
    except (ValueError,TypeError):err("422","Isian angka tidak valid")
def hp(p):return hashlib.pbkdf2_hmac("sha256",p.encode(),b"pejadin",50000).hex()
def err(code,msg,st=400):raise HTTPException(st,{"error":{"code":code,"message":msg}})
def sig(b):return hmac.new(SECRET.encode(),b.encode(),hashlib.sha256).hexdigest()
def touch(table:str, row_id:int):
    if not table or row_id is None: return
    ex(f"update {table} set updated_at=now() where id=?",(row_id,))
def user(req:Request):
    t=req.headers.get("authorization","")[7:] or req.query_params.get("token","")
    try:
        b,s=t.split(".");assert hmac.compare_digest(s,sig(b))
        d=json.loads(base64.urlsafe_b64decode(b));assert d["exp"]>time.time()
        u=q("select * from users where id=? and aktif=1",(d["id"],),True);assert u
    except Exception:err("401","Sesi tidak valid, silakan login ulang",401)
    return u
def need(*roles):
    def f(u=Depends(user)):
        if u["role"]!="ADMIN" and u["role"] not in roles:err("403","Peran tidak berhak",403)
        return u
    return f
NOW=lambda:dt.datetime.now().isoformat(" ","seconds")
def log(u,aksi,ent,eid,det=""):
    payload = json.dumps({"detail": det}) if det else json.dumps({})
    with pool.connection() as c:
        with c.transaction():
            c.execute(
                P("insert into audit(pengguna,aksi,entitas,eid,detail,waktu) values(?,?,?,?,?,?)"),
                (u["username"],aksi,ent,eid,det,NOW())
            )
            c.execute(
                P("insert into audit_log(user_id,aksi,entitas,entitas_id,detail,ip_address) values(?,?,?,?,?,?)"),
                (u.get("id"),aksi,ent,eid,payload,None)
            )

PROV="ACEH,SUMATERA UTARA,SUMATERA BARAT,RIAU,JAMBI,SUMATERA SELATAN,BENGKULU,LAMPUNG,KEPULAUAN BANGKA BELITUNG,KEPULAUAN RIAU,DKI JAKARTA,JAWA BARAT,JAWA TENGAH,DI YOGYAKARTA,JAWA TIMUR,BANTEN,BALI,NUSA TENGGARA BARAT,NUSA TENGGARA TIMUR,KALIMANTAN BARAT,KALIMANTAN TENGAH,KALIMANTAN SELATAN,KALIMANTAN TIMUR,KALIMANTAN UTARA,SULAWESI UTARA,SULAWESI TENGAH,SULAWESI SELATAN,SULAWESI TENGGARA,GORONTALO,SULAWESI BARAT,MALUKU,MALUKU UTARA,PAPUA,PAPUA BARAT,PAPUA SELATAN,PAPUA TENGAH,PAPUA PEGUNUNGAN,PAPUA BARAT DAYA".split(",")

def ensure_design_compatibility():
    with pool.connection() as c:
        def execute_compatibility_sql(statement,description):
            try:
                with c.transaction():
                    c.execute(statement)
            except psycopg.Error:
                logger.warning("Kompatibilitas skema dilewati untuk %s",description,exc_info=True)

        c.execute("create table if not exists audit_log(id bigserial primary key,user_id bigint,aksi text,entitas text,entitas_id bigint,detail jsonb,ip_address text,created_at timestamptz not null default now())")
        c.execute("create table if not exists dokumen_terbit(id bigserial primary key,surat_tugas_id bigint references st(id),spd_id bigint references spd(id),jenis_dokumen text,storage_key text,dicetak_oleh bigint,dicetak_pada timestamptz not null default now())")
        c.execute("create table if not exists uang_representasi(tahun int not null,kategori text not null,luar_kota_oh bigint not null,dalam_kota_8jam bigint not null,primary key(tahun,kategori))")
        c.execute("create table if not exists golongan_pangkat(golongan text primary key,nama_pangkat text not null,created_at timestamptz not null default now(),updated_at timestamptz not null default now())")
        for stmt in [
            "alter table users add column if not exists last_login timestamptz",
            "alter table users add column if not exists created_at timestamptz not null default now()",
            "alter table users add column if not exists updated_at timestamptz not null default now()",
            "alter table st add column if not exists created_at timestamptz not null default now()",
            "alter table st add column if not exists updated_at timestamptz not null default now()",
            "alter table st add column if not exists nama_kegiatan text",
            "alter table st add column if not exists tanggal_surat_tugas date",
            "alter table st add column if not exists tgl_tanda_tangan date",
            "alter table st add column if not exists nomor_bukti text",
            "alter table st add column if not exists uraian_kegiatan text",
            "alter table st add column if not exists keterangan_laporan text",
            "alter table st add column if not exists diterbitkan_pada timestamptz",
            "alter table spd add column if not exists created_at timestamptz not null default now()",
            "alter table spd add column if not exists updated_at timestamptz not null default now()",
            "alter table spd add column if not exists pangkat text",
            "alter table spd add column if not exists tingkat_biaya text",
            "alter table spd add column if not exists biaya_riil bigint",
            "alter table spd add column if not exists ppn bigint",
            "alter table spd add column if not exists pph bigint",
            "alter table spd add column if not exists tgl_tiba_pengesahan date",
            "alter table spd add column if not exists status text not null default 'DRAFT'",
            "alter table pegawai add column if not exists created_at timestamptz not null default now()",
            "alter table pegawai add column if not exists updated_at timestamptz not null default now()",
            "alter table pegawai add column if not exists tingkat_biaya text",
            "alter table pegawai add column if not exists pangkat text",
            "alter table tarif add column if not exists created_at timestamptz not null default now()",
            "alter table tarif add column if not exists updated_at timestamptz not null default now()",
            "alter table pejabat add column if not exists created_at timestamptz not null default now()",
            "alter table pejabat add column if not exists updated_at timestamptz not null default now()",
            "alter table dasar_hukum add column if not exists created_at timestamptz not null default now()",
            "alter table dasar_hukum add column if not exists updated_at timestamptz not null default now()",
            "alter table satker add column if not exists created_at timestamptz not null default now()",
            "alter table satker add column if not exists updated_at timestamptz not null default now()",
            "alter table uang_representasi add column if not exists created_at timestamptz not null default now()",
            "alter table uang_representasi add column if not exists updated_at timestamptz not null default now()",
            "alter table tarif add column if not exists uang_harian_dalam_kota_8jam bigint not null default 0",
            "alter table tarif add column if not exists uang_harian_diklat bigint not null default 0",
        ]:
            c.execute(stmt)
        c.execute("alter table uang_representasi enable row level security")
        c.execute("alter table golongan_pangkat enable row level security")
        c.execute("insert into satker(k,v) values('nama_satker_panjang',''),('dipa_raw',''),('klasifikasi_anggaran',''),('rekening_bendahara',''),('kota_kedudukan',''),('tanggal_dipa',''),('kode_akun_belanja','') on conflict(k) do nothing")
        migrate_legacy_satker(c)
        c.execute("alter table satker_tahunan enable row level security")
        try:
            with c.transaction():
                orphaned_spd = ensure_spd_pegawai_foreign_key(c)
            if orphaned_spd:
                logger.warning(
                    "Foreign key SPD-pegawai dipasang tanpa memvalidasi baris lama; ditemukan %s baris SPD tanpa pegawai terkait",
                    orphaned_spd,
                )
        except psycopg.Error:
            logger.warning(
                "Foreign key SPD-pegawai gagal dipasang",
                exc_info=True,
            )
        for column in ("tgl_berangkat","tgl_kembali","tgl_tiba"):
            data_type=c.execute(
                "select data_type from information_schema.columns where table_schema=current_schema() and table_name='st' and column_name=%s",
                (column,)
            ).fetchone()
            if data_type and data_type["data_type"]!="date":
                execute_compatibility_sql(
                    f"alter table st alter column {column} type date using case when {column} is null then null else {column}::date end",
                    f"konversi kolom {column}"
                )
        try:
            with c.transaction():
                c.execute("create or replace function set_updated_at() returns trigger as $$ begin new.updated_at = now(); return new; end; $$ language plpgsql")
                for tbl in ["users","pegawai","tarif","pejabat","dasar_hukum","satker","satker_tahunan","st","spd","golongan_pangkat"]:
                    c.execute(f"drop trigger if exists trg_{tbl}_updated_at on {tbl}")
                    c.execute(f"create trigger trg_{tbl}_updated_at before update on {tbl} for each row execute procedure set_updated_at()")
        except psycopg.Error:
            logger.exception("Gagal memasang trigger updated_at untuk kompatibilitas desain")
        execute_compatibility_sql("create or replace view surat_tugas as select * from st","view surat_tugas")
        execute_compatibility_sql("create or replace view audit_log_view as select * from audit_log","view audit_log_view")

@app.on_event("startup")
def init():
    try:
        pool.open(wait=True,timeout=30)
    except Exception as e:
        raise RuntimeError("Gagal konek ke database Supabase. Periksa DATABASE_URL di backend/.env (pakai Session pooler & password benar). Detail: "+str(e))
    if q("select count(*) n from information_schema.tables where table_schema='public' and table_name='users'",one=True)["n"]:
        ensure_design_compatibility()
        return
    with pool.connection() as c:
        c.execute(open(os.path.join(D,"..","database","schema.sql"),encoding="utf-8").read())
        E=lambda s,a=():c.execute(P(s),a)
        for u,r,n,p in [("admin","ADMIN","Administrator",None),("operator","OPERATOR","Operator Satker",None),("ppk","PPK","Dedi Adrianto, SE, MM",None),("bendahara","BENDAHARA","Bendahara Pengeluaran",None),("pegawai","PEGAWAI","Pegawai Contoh",1)]:
            E("insert into users(username,password_hash,nama,role,pegawai_id) values(?,?,?,?,?)",(u,hp(u+"123"),n,r,p))
        for r in [("Arnold A.P. Ritiauw","196903102026211001","IV/d","Sekretaris Direktorat Jenderal","I.a","Sekretariat Ditjen"),("Budi Contoh","197001012000031001","IV/c","Kepala Bagian Umum","II.a","Bagian Umum"),("Siti Rahma","198002022005012002","III/d","Kepala Subbagian","III.a","Bagian Keuangan"),("Andi Wijaya","198505052010011003","III/b","Pelaksana","IV.a","Bagian Umum"),("Dewi Lestari","199001012015022001","III/a","Pelaksana","IV.a","Bagian Program"),("Rudi Hartono","198803032012011004","III/c","Analis Anggaran","III.a","Bagian Keuangan")]:
            E("insert into pegawai(nama,nip,golongan,jabatan,eselon,unit_kerja) values(?,?,?,?,?,?)",r)
        for i,p in enumerate(PROV):
            h,i1=(440000,4682000) if p=="NUSA TENGGARA BARAT" else (370000+(i%6)*30000,3500000+(i%7)*250000)
            E("insert into tarif(tahun,provinsi,harian,inap1,inap2,inap3,inap4,taksi) values(2026,?,?,?,?,?,?,?)",(p,h,i1,int(i1*.75),int(i1*.5),int(i1*.4),150000+(i%4)*50000))
        for r,n in [("PPK","Dedi Adrianto, SE, MM"),("BENDAHARA","Bendahara Pengeluaran"),("SEKDITJEN","Arnold A.P. Ritiauw"),("KABAG","Budi Contoh"),("PELAKSANA_ADM","Siti Rahma"),("PEMBUAT_DAFTAR","Andi Wijaya"),("PENGESAH","Arnold A.P. Ritiauw")]:
            E("insert into pejabat(peran,nama,nip,jabatan) values(?,?,?,?)",(r,n,"000000000000000000",r.title()))
        for i,x in enumerate(["Peraturan Menteri Keuangan Nomor 113/PMK.05/2012 tentang Perjalanan Dinas Dalam Negeri","Peraturan Menteri Keuangan Nomor 32 Tahun 2025 tentang Standar Biaya Masukan TA 2026","DIPA Satuan Kerja Tahun Anggaran 2026"]):E("insert into dasar_hukum(urutan,teks) values(?,?)",(i+1,x))
        for k,v in {"nama_kementerian":"KEMENTERIAN PEKERJAAN UMUM","nama_ditjen":"DIREKTORAT JENDERAL SUMBER DAYA AIR","nama_satker":"Sekretariat Direktorat Jenderal Sumber Daya Air","nama_satker_panjang":"","alamat":"Jakarta","kode_satker":"694257","nomor_dipa":"","tanggal_dipa":"","dipa_raw":"","klasifikasi_anggaran":"","kode_akun_belanja":"","rekening_bendahara":"","nomor_sptjb":"","akun_default":"7755.EBA.962.182.524111 D","kota":"Jakarta","kota_kedudukan":"Jakarta"}.items():E("insert into satker(k,v) values(?,?)",(k,v))
    ensure_design_compatibility()

@app.post("/api/login")
async def login(r:Request):
    b=await r.json();u=q("select * from users where username=? and aktif=1",(b.get("username"),),True)
    if not u or u["password_hash"]!=hp(b.get("password","")):err("401","Username atau password salah",401)
    ex("update users set last_login=?,updated_at=now() where id=?",(NOW(),u["id"]))
    u["last_login"]=NOW();u["updated_at"]=NOW()
    t=base64.urlsafe_b64encode(json.dumps({"id":u["id"],"exp":time.time()+28800}).encode()).decode()
    log(u,"LOGIN","users",u["id"]);u.pop("password_hash");return {"token":t+"."+sig(t),"user":u}
@app.post("/api/v1/auth/login")
async def login_v1(r:Request):
    return await login(r)
@app.get("/api/me")
def me(u=Depends(user)):u.pop("password_hash");return u
@app.get("/api/v1/auth/me")
def me_v1(u=Depends(user)):u.pop("password_hash");return u
T={"pegawai":("pegawai","nama nip golongan pangkat jabatan eselon tingkat_biaya unit_kerja aktif"),"tarif":("tarif","tahun provinsi harian inap1 inap2 inap3 inap4 taksi uang_harian_dalam_kota_8jam uang_harian_diklat"),"pejabat":("pejabat","peran nama nip jabatan aktif"),"dasar_hukum":("dasar_hukum","urutan teks aktif"),"users":("users","username nama role pegawai_id aktif")}
@app.get("/api/ref/golongan-pangkat",dependencies=[Depends(user)])
def golongan_pangkat():return q("select golongan,nama_pangkat from golongan_pangkat order by golongan")
@app.get("/api/v1/ref/golongan-pangkat",dependencies=[Depends(user)])
def golongan_pangkat_v1():return golongan_pangkat()
@app.get("/api/ref/satker")
def gs(u=Depends(user)):return {r["k"]:r["v"] for r in q("select * from satker")}
@app.get("/api/v1/ref/satker")
def gs_v1(u=Depends(user)):return gs(u)
@app.put("/api/ref/satker")
async def ps(r:Request,u=Depends(need())):
    values=(await r.json())
    if "tanggal_dipa" in values and values["tanggal_dipa"] not in (None,""):
        try:
            values["tanggal_dipa"]=dt.date.fromisoformat(values["tanggal_dipa"])
        except (TypeError,ValueError):
            err("422","Tanggal DIPA harus berformat YYYY-MM-DD",422)
    with pool.connection() as c:
        with c.transaction():
            for k,v in values.items():
                c.execute(P("insert into satker(k,v) values(?,?) on conflict (k) do update set v=excluded.v"),(k,str(v)))
            sync_legacy_satker(c,values)
    log(u,"UPDATE","satker",0);return {"ok":1}
@app.put("/api/v1/ref/satker")
async def ps_v1(r:Request,u=Depends(need())):
    return await ps(r,u)
@app.get("/api/ref/{n}")
def lst(n:str,u=Depends(user)):
    if n not in T:err("404","Tidak ada",404)
    if n=="users" and u["role"]!="ADMIN":err("403","Hanya Admin",403)
    return q(f"select id,{T[n][1].replace(' ',',')} from {T[n][0]} order by "+("tahun desc,provinsi" if n=="tarif" else "id"))
@app.get("/api/v1/ref/{n}")
def lst_v1(n:str,u=Depends(user)):
    return lst(n,u)
@app.post("/api/ref/{n}")
async def add(n:str,r:Request,u=Depends(need())):
    if n not in T:err("404","Tidak ada",404)
    b=await r.json();t,cs=T[n];cs=cs.split();b=ci({k:v for k,v in b.items() if k in cs or k=="password"})
    if n=="users":b["password_hash"]=hp(b.pop("password","") or "ganti123")
    if n=="pegawai" and not b.get("nip"):err("422","NIP wajib")
    oc=""
    if n=="tarif":
        if not(b.get("tahun") and b.get("provinsi")):err("422","Tahun dan provinsi wajib diisi")
        cu=[k for k in b if k not in("tahun","provinsi")];oc=" on conflict (tahun,provinsi) do "+("update set "+",".join(f"{k}=excluded.{k}" for k in cu) if cu else "nothing")
    try:i=ex(f"insert into {t}({','.join(b)}) values({','.join('?'*len(b))}){oc}",list(b.values()),True)
    except psycopg.errors.IntegrityError:err("409","Data duplikat (mis. NIP/username sudah ada) atau tidak valid",409)
    touch(t, i)
    log(u,"CREATE",n,i,json.dumps(b)[:200]);return {"id":i}
@app.post("/api/v1/ref/{n}")
async def add_v1(n:str,r:Request,u=Depends(need())):
    return await add(n,r,u)
@app.put("/api/ref/{n}/{i}")
async def upd(n:str,i:int,r:Request,u=Depends(need())):
    b=await r.json();t,cs=T[n];b=ci({k:v for k,v in b.items() if k in cs.split()})
    ex(f"update {t} set "+",".join(f"{k}=?" for k in b)+" where id=?",list(b.values())+[i]);touch(t,i);log(u,"UPDATE",n,i,json.dumps(b)[:200]);return {"ok":1}
@app.put("/api/v1/ref/{n}/{i}")
async def upd_v1(n:str,i:int,r:Request,u=Depends(need())):
    return await upd(n,i,r,u)
def dstr(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return str(v)

def rep_of(e):return 200000 if e.startswith("I.") else 150000 if e.startswith("II.") else 0
def hitung(s,pg,hari,malam,tinap,tiket,lain):
    if not pg["eselon"] or not pg["golongan"]:
        err("V7",f"Eselon & golongan '{pg['nama']}' belum terisi")
    if hari<1 or malam<0 or min(tinap,tiket,lain)<0:
        err("V4","Hari harus minimal 1 dan biaya/malam inap tidak boleh negatif")
    tanggal_berangkat = dt.date.fromisoformat(dstr(s["tgl_berangkat"]))
    tanggal_kembali = dt.date.fromisoformat(dstr(s["tgl_kembali"]))
    tarif_db=q(
        "select * from tarif where tahun=? and provinsi=?",
        (tanggal_berangkat.year,s["provinsi"]),True
    )
    if not tarif_db:
        err("V3",f"Tarif {s['provinsi']} tahun {tanggal_berangkat.year} belum tersedia")
    tarif_required=("harian","inap1","inap2","inap3","inap4","taksi")
    if any(tarif_db.get(k) is None or int(tarif_db[k])<0 for k in tarif_required):
        err("V3",f"Tarif SBM {s['provinsi']} tahun {tanggal_berangkat.year} belum lengkap")
    pegawai=PegawaiHitung(
        id=str(pg["id"]),nama=pg["nama"],nip_nrp=pg["nip"],golongan=pg["golongan"],
        pangkat=pg.get("pangkat") or "",jabatan=pg["jabatan"] or "",
        eselon=pg["eselon"],tingkat_biaya=pg.get("tingkat_biaya")
    )
    st=SuratTugasHitung(
        nomor_surat_tugas=s.get("nomor") or "",tanggal_surat_tugas=s.get("tanggal_surat_tugas") or tanggal_berangkat,
        maksud_perjalanan=s.get("maksud") or "",asal=s.get("asal") or "",
        tujuan=s.get("tujuan") or "",provinsi_tarif=s["provinsi"],
        tgl_berangkat=tanggal_berangkat,tgl_kembali=tanggal_kembali,
        tgl_tanda_tangan=s.get("tgl_tanda_tangan") or tanggal_berangkat,
        akun_anggaran=s.get("akun") or "",nama_kegiatan=s.get("nama_kegiatan"),
        alat_angkutan=s.get("alat"),jenis_kuitansi=s.get("jenis"),status=s.get("status","DRAFT")
    )
    tarif=TarifSBM(
        tahun=int(tarif_db["tahun"]),provinsi=tarif_db["provinsi"],
        uang_harian_luar_kota=int(tarif_db["harian"] or 0),
        uang_harian_dalam_kota_8jam=int(tarif_db.get("uang_harian_dalam_kota_8jam") or 0),
        uang_harian_diklat=int(tarif_db.get("uang_harian_diklat") or 0),
        inap_pejabat_negara_eselon1=int(tarif_db["inap1"] or 0),
        inap_eselon2=int(tarif_db["inap2"] or 0),
        inap_eselon3_gol4=int(tarif_db["inap3"] or 0),
        inap_eselon4_gol321=int(tarif_db["inap4"] or 0),
        taksi=int(tarif_db["taksi"] or 0)
    )
    rep_rates={
        (int(r["tahun"]),r["kategori"]):UangRepresentasiRate(
            tahun=int(r["tahun"]),kategori=r["kategori"],
            luar_kota_oh=int(r["luar_kota_oh"]),dalam_kota_8jam=int(r["dalam_kota_8jam"])
        ) for r in q("select * from uang_representasi where tahun=?",(tanggal_berangkat.year,))
    }
    try:
        hasil=hitung_spd(
            st,SPDMasukan(pegawai,hari,malam,tinap,tiket,lain),
            {(tarif.tahun,tarif.provinsi.upper()):tarif},rep_rates,tanggal_berangkat.year
        )
    except ErroValidasiHitung as exc:
        message=str(exc)
        code=message.split(":",1)[0] if ":" in message else "422"
        err(code,message)
    return dict(
        tarif_harian=hasil.tarif_uang_harian,uang_harian=hasil.uang_harian,
        uang_inap=hasil.uang_penginapan,taksi=hasil.taksi,
        tarif_rep=hasil.tarif_representasi,rep=hasil.uang_representasi,
        jumlah=hasil.jumlah,batas=hasil.batas_sbm_inap,
        flag=int(hasil.flag_lebih_sbm)
    )

def snapshot_spd(pg,hasil):
    return {
        "nama":pg["nama"],"nip":pg["nip"],"golongan":pg["golongan"],
        "pangkat":pg.get("pangkat"),"tingkat_biaya":pg.get("tingkat_biaya"),
        "jabatan":pg["jabatan"],"eselon":pg["eselon"],
        **hasil
    }

def lama(s):
    tgl_berangkat = dstr(s["tgl_berangkat"])
    tgl_kembali = dstr(s["tgl_kembali"])
    return (dt.date.fromisoformat(tgl_kembali)-dt.date.fromisoformat(tgl_berangkat)).days+1
@app.get("/api/st")
def stl(status:str="",cari:str="",u=Depends(user)):
    w,a=["1=1"],[]
    stats_filter=""
    stats_args=[]
    if u["role"]=="PEGAWAI":
        stats_filter=" and pegawai_id=?"
        stats_args=[u["pegawai_id"]]*3
    if status:w.append("s.status=?");a.append(status)
    if cari:w.append("(s.nomor ilike ? or s.maksud ilike ? or s.tujuan ilike ?)");a+=[f"%{cari}%"]*3
    if u["role"]=="PEGAWAI":w.append("exists(select 1 from spd where st_id=s.id and pegawai_id=?)");a.append(u["pegawai_id"])
    return q(
        "select s.*,"
        f"(select count(*) from spd where st_id=s.id{stats_filter}) jml,"
        f"cast(coalesce((select sum(jumlah) from spd where st_id=s.id{stats_filter}),0) as bigint) total,"
        f"coalesce((select max(flag) from spd where st_id=s.id{stats_filter}),0) flag "
        "from st s where "+" and ".join(w)+" order by s.id desc",
        stats_args+a
    )
@app.get("/api/v1/st")
def stl_v1(status:str="",cari:str="",u=Depends(user)):
    return stl(status,cari,u)
@app.post("/api/st")
async def stc(r:Request,u=Depends(need("OPERATOR"))):
    b=await r.json()
    if not b.get("tgl_berangkat") or not b.get("tgl_kembali"):err("422","Tanggal wajib diisi")
    try:
        tgl_berangkat=dt.date.fromisoformat(b["tgl_berangkat"])
        tgl_kembali=dt.date.fromisoformat(b["tgl_kembali"])
    except (TypeError,ValueError):err("V1","Format tanggal harus YYYY-MM-DD")
    if tgl_kembali<tgl_berangkat:err("V1","Tanggal kembali tidak boleh sebelum tanggal berangkat")
    if not q("select 1 from tarif where provinsi=?",(b.get("provinsi"),)):err("V2","Provinsi tarif wajib dipilih dari daftar")
    k="maksud asal tujuan provinsi tgl_berangkat tgl_kembali alat akun jenis".split()
    extra=("nama_kegiatan","tanggal_surat_tugas","tgl_tanda_tangan","nomor_bukti","uraian_kegiatan","keterangan_laporan")
    for field in extra:
        if field in b:
            if field in("tanggal_surat_tugas","tgl_tanda_tangan"):
                if b[field]:
                    try:b[field]=dt.date.fromisoformat(b[field])
                    except (TypeError,ValueError):err("422",f"Format {field} harus YYYY-MM-DD")
                else:
                    b[field]=None
            k.append(field)
    v=[b.get(x) for x in k]
    i=ex(f"insert into st({','.join(k)},tahun,dibuat_oleh) values({','.join('?'*(len(k)+2))})",v+[tgl_berangkat.year,u["username"]],True);touch("st",i);log(u,"CREATE","st",i);return {"id":i,"status":"DRAFT"}
@app.post("/api/v1/st")
async def stc_v1(r:Request,u=Depends(need("OPERATOR"))):
    return await stc(r,u)
@app.get("/api/st/{i}")
def std(i:int,u=Depends(user)):
    s=q("select * from st where id=?",(i,),True)
    if not s:err("404","Tidak ditemukan",404)
    s["spd"]=filter_spd_for_user(q("select * from spd where st_id=? order by urutan",(i,)),u)
    if u["role"]=="PEGAWAI" and not s["spd"]:err("403","Bukan SPD Anda",403)
    s["lama"]=lama(s);s["total"]=sum(x["jumlah"] for x in s["spd"]);return s
@app.get("/api/v1/st/{i}")
def std_v1(i:int,u=Depends(user)):
    return std(i,u)
@app.put("/api/st/{i}/pdf-metadata")
@app.put("/api/v1/st/{i}/pdf-metadata")
async def update_pdf_metadata(i:int,r:Request,u=Depends(need("PPK"))):
    body=await r.json()
    if not isinstance(body,dict) or set(body)-{"tanggal_surat_tugas","tgl_tanda_tangan","spd"}:
        err("422","Isian koreksi PDF tidak valid",422)
    dates={}
    for field in ("tanggal_surat_tugas","tgl_tanda_tangan"):
        if field in body:
            try:
                dates[field]=dt.date.fromisoformat(body[field])
            except (TypeError,ValueError):
                err("422",f"Format {field} harus YYYY-MM-DD",422)
    spd_updates=body.get("spd",[])
    if not isinstance(spd_updates,list):
        err("422","Data koreksi pelaksana harus berupa daftar",422)
    changes=[]
    with pool.connection() as c:
        st=c.execute("select id,status,tanggal_surat_tugas,tgl_tanda_tangan from st where id=%s for update",(i,)).fetchone()
        if not st:
            err("404","Tidak ditemukan",404)
        if st["status"] not in("TERBIT","SELESAI"):
            err("409","Koreksi PDF hanya tersedia untuk ST TERBIT/SELESAI",409)
        for field,value in dates.items():
            if st[field] and st[field]!=value:
                err("409",f"{field} sudah terisi dan tidak dapat ditimpa",409)
            if not st[field]:
                c.execute(P(f"update st set {field}=? where id=?"),(value,i))
                changes.append(field)
        for item in spd_updates:
            if not isinstance(item,dict) or set(item)-{"id","pangkat","tingkat_biaya"} or "id" not in item:
                err("422","Data koreksi snapshot pelaksana tidak valid",422)
            spd_id=item["id"]
            if isinstance(spd_id,bool) or not isinstance(spd_id,int):
                err("422","ID snapshot pelaksana tidak valid",422)
            row=c.execute("select id,pangkat,tingkat_biaya from spd where id=%s and st_id=%s for update",(spd_id,i)).fetchone()
            if not row:
                err("404","Snapshot pelaksana tidak ditemukan pada ST ini",404)
            for field in ("pangkat","tingkat_biaya"):
                if field not in item:
                    continue
                value=item[field]
                if not isinstance(value,str) or not value.strip() or len(value.strip())>120:
                    err("422",f"{field} wajib diisi (maksimal 120 karakter)",422)
                value=value.strip()
                if row[field] and row[field]!=value:
                    err("409",f"{field} snapshot sudah terisi dan tidak dapat ditimpa",409)
                if not row[field]:
                    c.execute(P(f"update spd set {field}=? where id=?"),(value,spd_id))
                    changes.append(f"spd:{spd_id}:{field}")
        if not changes:
            err("422","Tidak ada kolom kosong yang dikoreksi",422)
        c.execute("update st set updated_at=now() where id=%s",(i,))
        detail="Koreksi metadata PDF: "+", ".join(changes)
        c.execute(
            "insert into audit(pengguna,aksi,entitas,eid,detail,waktu) values(%s,%s,%s,%s,%s,%s)",
            (u["username"],"PDF_METADATA_BACKFILL","st",i,detail,NOW())
        )
        c.execute(
            "insert into audit_log(user_id,aksi,entitas,entitas_id,detail,ip_address) values(%s,%s,%s,%s,%s,%s)",
            (u.get("id"),"PDF_METADATA_BACKFILL","st",i,json.dumps({"detail":detail}),None)
        )
    return {"id":i,"updated":changes}
@app.delete("/api/st/{i}")
def stx(i:int,u=Depends(need("OPERATOR"))):
    if std(i,u)["status"]!="DRAFT":err("409","Hanya Draft yang bisa dihapus",409)
    ex("delete from spd where st_id=?",(i,));ex("delete from st where id=?",(i,));log(u,"DELETE","st",i);return {"ok":1}
@app.delete("/api/v1/st/{i}")
def stx_v1(i:int,u=Depends(need("OPERATOR"))):
    return stx(i,u)
@app.post("/api/st/{i}/spd")
async def spa(i:int,r:Request,u=Depends(need("OPERATOR"))):
    s=std(i,u);b=await r.json()
    if s["status"]!="DRAFT":err("409","Hanya Draft yang bisa diubah",409)
    pg=q("select * from pegawai where id=? and aktif=1",(b.get("pegawai_id"),),True)
    if not pg:err("404","Pegawai tidak ditemukan",404)
    if any(x["pegawai_id"]==pg["id"] for x in s["spd"]):err("409","Pegawai sudah ada di surat tugas ini",409)
    o=q("select s.id from spd p join st s on s.id=p.st_id where p.pegawai_id=? and s.id!=? and s.status!='DIBATALKAN' and s.tgl_berangkat<=? and s.tgl_kembali>=?",(pg["id"],i,s["tgl_kembali"],s["tgl_berangkat"]))
    if o:err("V5",f"'{pg['nama']}' sudah punya perjalanan lain di tanggal yang bertumpuk (ST #{o[0]['id']})")
    try:
        n=lambda k:int(b.get(k) or 0)
        hari=n("hari") or s["lama"];malam=n("malam");tarif_inap=n("tarif_inap")
        tiket=n("tiket");lain=n("lain")
        biaya_riil=int(b["biaya_riil"]) if b.get("biaya_riil") not in(None,"") else None
        ppn=int(b["ppn"]) if b.get("ppn") not in(None,"") else None
        pph=int(b["pph"]) if b.get("pph") not in(None,"") else None
    except (TypeError,ValueError):err("422","Jumlah hari dan biaya harus berupa bilangan bulat")
    if min(tiket,lain,tarif_inap)<0 or any(x is not None and x<0 for x in (biaya_riil,ppn,pph)):err("V4","Biaya tidak boleh negatif")
    h=hitung(s,pg,hari,malam,tarif_inap,tiket,lain)
    d=dict(st_id=i,urutan=len(s["spd"])+1,pegawai_id=pg["id"],nama=pg["nama"],nip=pg["nip"],golongan=pg["golongan"],pangkat=pg.get("pangkat"),tingkat_biaya=pg.get("tingkat_biaya"),jabatan=pg["jabatan"],eselon=pg["eselon"],hari=hari,malam=malam,tarif_inap=tarif_inap,tiket=tiket,lain=lain,biaya_riil=biaya_riil,ppn=ppn,pph=pph,**h)
    d["id"]=ex(f"insert into spd({','.join(d)}) values({','.join('?'*len(d))})",list(d.values()),True);touch("spd",d["id"]);log(u,"CREATE","spd",d["id"]);return d
@app.post("/api/v1/st/{i}/spd")
async def spa_v1(i:int,r:Request,u=Depends(need("OPERATOR"))):
    return await spa(i,r,u)
@app.delete("/api/st/{i}/spd/{p}")
def spx(i:int,p:int,u=Depends(need("OPERATOR"))):
    if std(i,u)["status"]!="DRAFT":err("409","Hanya Draft yang bisa diubah",409)
    ex("delete from spd where id=? and st_id=?",(p,i));log(u,"DELETE","spd",p);return {"ok":1}
@app.delete("/api/v1/st/{i}/spd/{p}")
def spx_v1(i:int,p:int,u=Depends(need("OPERATOR"))):
    return spx(i,p,u)
@app.post("/api/st/{i}/{act}")
async def aksi(i:int,act:str,r:Request,u=Depends(user)):
    raw=await r.body();b=json.loads(raw) if raw else {};s=std(i,u);st=s["status"]
    M={"ajukan":("OPERATOR","DRAFT","DIAJUKAN"),"kembalikan":("PPK","DIAJUKAN","DRAFT"),"terbitkan":("PPK","DIAJUKAN","TERBIT"),"selesaikan":("OPERATOR","TERBIT","SELESAI"),"batalkan":("PPK",None,"DIBATALKAN")}
    if act not in M:err("404","Aksi tidak dikenal",404)
    rl,fr,to=M[act]
    if u["role"] not in("ADMIN",rl) and not(act=="selesaikan" and u["role"]=="PPK"):err("403","Peran tidak berhak",403)
    if (fr and st!=fr) or st in("SELESAI","DIBATALKAN"):err("409",f"Status {st} tidak bisa {act}",409)
    if act=="ajukan" and not s["spd"]:err("422","Tambahkan minimal 1 pelaksana")
    if act in("kembalikan","batalkan") and not b.get("alasan"):err("422","Alasan wajib diisi")
    if act=="terbitkan":
        nm=(b.get("nomor") or "").strip()
        if not nm:err("422","Nomor surat tugas wajib diisi")
        try:
            with pool.connection() as c:
                locked_st=c.execute("select * from st where id=%s for update",(i,)).fetchone()
                if not locked_st:err("404","Tidak ditemukan",404)
                if locked_st["status"]!="DIAJUKAN":
                    err("409",f"Status {locked_st['status']} tidak bisa terbitkan",409)
                duplicate=c.execute(
                    "select 1 from st where nomor=%s and tahun=%s and id!=%s",
                    (nm,locked_st["tahun"],i)
                ).fetchone()
                if duplicate:err("V6",f"Nomor Surat Tugas '{nm}' sudah dipakai.")
                spd_rows=c.execute(
                    "select * from spd where st_id=%s order by urutan for update",(i,)
                ).fetchall()
                if not spd_rows:err("422","Tambahkan minimal 1 pelaksana")
                for row in spd_rows:
                    pg=c.execute(
                        "select * from pegawai where id=%s and aktif=1 for share",
                        (row["pegawai_id"],)
                    ).fetchone()
                    if not pg:err("V7","Pegawai pelaksana tidak ditemukan atau tidak aktif")
                    hasil=hitung(
                        locked_st,pg,row["hari"],row["malam"],row["tarif_inap"],
                        row["tiket"],row["lain"]
                    )
                    snapshot=snapshot_spd(pg,hasil)
                    assignments=",".join(f"{column}=%s" for column in snapshot)
                    c.execute(
                        f"update spd set {assignments},status='TERBIT',updated_at=now() "
                        "where id=%s and st_id=%s",
                        (*snapshot.values(),row["id"],i)
                    )
                c.execute(
                    "update st set nomor=%s,diterbitkan_oleh=%s,diterbitkan_pada=now(),"
                    "status='TERBIT',alasan=NULL,updated_at=now() where id=%s",
                    (nm,u["nama"],i)
                )
        except psycopg.errors.UniqueViolation:
            err("V6",f"Nomor Surat Tugas '{nm}' sudah dipakai.")
        log(u,"TERBITKAN","st",i,nm)
        return {"id":i,"status":"TERBIT"}
    if act=="selesaikan":ex("update st set tgl_tiba=? where id=?",(b.get("tgl_tiba") or s["tgl_kembali"],i))
    ex("update st set status=?,alasan=?,updated_at=now() where id=?",(to,b.get("alasan"),i));touch("st",i);log(u,act.upper(),"st",i,b.get("alasan") or b.get("nomor") or "");return {"id":i,"status":to}
@app.post("/api/v1/st/{i}/{act}")
async def aksi_v1(i:int,act:str,r:Request,u=Depends(user)):
    return await aksi(i,act,r,u)
@app.get("/api/audit")
def aud(u=Depends(need())):return q("select * from audit order by id desc limit 300")
@app.get("/api/v1/audit")
def aud_v1(u=Depends(need())):return aud(u)
@app.get("/api/admin/audit-log")
@app.get("/api/v1/admin/audit-log")
def audit_log(entitas:str="",entitas_id:int=None,user_id:int=None,u=Depends(need())):
    w,a=["1=1"],[]
    if entitas:w.append("entitas=?");a.append(entitas)
    if entitas_id is not None:w.append("entitas_id=?");a.append(entitas_id)
    if user_id is not None:w.append("user_id=?");a.append(user_id)
    return q("select * from audit_log where "+" and ".join(w)+" order by id desc limit 300",a)
def rp(n):
    amount=int(n or 0)
    return ("-Rp " if amount<0 else "Rp ")+f"{abs(amount):,}".replace(",",".")
def html_escape(value):
    return (str(value) if value is not None else "").replace("&","&amp;").replace("<","&lt;").replace(">","&gt;").replace('"',"&quot;").replace("'","&#x27;")
def tb(n):
    s=["","satu","dua","tiga","empat","lima","enam","tujuh","delapan","sembilan","sepuluh","sebelas"];j=lambda x:(" "+tb(x)) if x else ""
    if n<12:return s[n]
    if n<20:return tb(n-10)+" belas"
    if n<100:return tb(n//10)+" puluh"+j(n%10)
    if n<200:return "seratus"+j(n-100)
    if n<1000:return tb(n//100)+" ratus"+j(n%100)
    if n<2000:return "seribu"+j(n-1000)
    if n<10**6:return tb(n//1000)+" ribu"+j(n%1000)
    if n<10**9:return tb(n//10**6)+" juta"+j(n%10**6)
    return tb(n//10**9)+" miliar"+j(n%10**9)
@app.get("/api/doc/{i}/{jenis}",response_class=HTMLResponse)
def doc(i:int,jenis:str,u=Depends(user)):
    s=std(i,u);S=gs(u);pj={p["peran"]:p for p in q("select * from pejabat where aktif=1")}
    if s["status"] not in("TERBIT","SELESAI"):err("V8","Dokumen hanya bisa dicetak bila status TERBIT/SELESAI",403)
    if u["role"]=="PEGAWAI":
        if jenis!="spd":err("403","Pegawai hanya dapat mencetak SPD miliknya",403)
        s["spd"]=[x for x in s["spd"] if x["pegawai_id"]==u["pegawai_id"]]
        if not s["spd"]:err("403","Bukan SPD Anda",403)
    log(u,"CETAK","st",i,jenis)
    ex("insert into dokumen_terbit(surat_tugas_id,spd_id,jenis_dokumen,dicetak_oleh) values(?,?,?,?)",(i,None,jenis,u["id"]))
    kop=f"<div style='text-align:center'><b>{html_escape(S['nama_kementerian'])}<br>{html_escape(S['nama_ditjen'])}</b><br>{html_escape(S['nama_satker'])}<br><small>{html_escape(S['alamat'])}</small><hr></div>"
    ttd=lambda r:f"<div style='float:right;text-align:center;margin-top:30px'>{html_escape(S['kota'])}, {html_escape(s['tgl_berangkat'])}<br>{html_escape(pj[r]['jabatan'])}<br><br><br><br><u>{html_escape(pj[r]['nama'])}</u><br>NIP {html_escape(pj[r]['nip'])}</div><div style='clear:both'></div>" if r in pj else ""
    ident=f"<p>Nomor: <b>{html_escape(s['nomor'])}</b><br>Maksud: {html_escape(s['maksud'])}<br>Tujuan: {html_escape(s['tujuan'])} ({html_escape(s['provinsi'])})<br>Tanggal: {html_escape(s['tgl_berangkat'])} s/d {html_escape(s['tgl_kembali'])} ({html_escape(s['lama'])} hari)</p>"
    def tbl(rows,h):
        header="<tr>"+"".join(f"<th>{html_escape(x)}</th>" for x in h)+"</tr>"
        body=[]
        for row in rows:
            cells=[]
            for value in row:
                escaped=html_escape(value)
                cells.append(f"<td><b>{escaped}</b></td>" if row[0]=="Jumlah" else f"<td>{escaped}</td>")
            body.append("<tr>"+"".join(cells)+"</tr>")
        return "<table border=1 cellpadding=5 style='border-collapse:collapse;width:100%'>"+header+"".join(body)+"</table>"
    sp=s["spd"];tot=sum(x["jumlah"] for x in sp)
    if jenis=="st":b=f"<h3 align=center>SURAT TUGAS<br><small>Nomor: {html_escape(s['nomor'])}</small></h3><p>Dasar:</p><ol>"+"".join(f"<li>{html_escape(d['teks'])}</li>" for d in q("select * from dasar_hukum where aktif=1 order by urutan"))+"</ol><p>Menugaskan:</p>"+tbl([(x["urutan"],x["nama"],x["nip"],x["jabatan"]) for x in sp],["No","Nama","NIP","Jabatan"])+f"<p>Untuk {html_escape(s['maksud'])} di {html_escape(s['tujuan'])} pada {html_escape(s['tgl_berangkat'])} s/d {html_escape(s['tgl_kembali'])}.</p>"+ttd("SEKDITJEN")
    elif jenis=="spd":b="".join(f"<div style='page-break-after:always'><h3 align=center>SURAT PERJALANAN DINAS (SPD)</h3>{ident}<p>Pelaksana: <b>{html_escape(x['nama'])}</b> / {html_escape(x['nip'])} / {html_escape(x['golongan'])} / {html_escape(x['jabatan'])}<br>Alat angkut: {html_escape(s['alat'])}<br>Akun: {html_escape(s['akun'])}</p>{ttd('PPK')}</div>" for x in sp)
    elif jenis=="rincian":b="".join(f"<div style='page-break-after:always'><h3 align=center>RINCIAN BIAYA PERJALANAN DINAS</h3>{ident}<p>{html_escape(x['nama'])}</p>"+tbl([("Uang harian",f"{x['hari']} x {rp(x['tarif_harian'])}",rp(x['uang_harian'])),("Penginapan (at-cost)",f"{x['malam']} x {rp(x['tarif_inap'])}",rp(x['uang_inap'])),("Transport/tiket","",rp(x['tiket'])),("Taksi","",rp(x['taksi'])),("Representasi",f"{x['hari']} x {rp(x['tarif_rep'])}",rp(x['rep'])),("Biaya lain","",rp(x['lain'])),("Jumlah","",rp(x['jumlah']))],["Uraian","Perhitungan","Jumlah"])+f"<p><i>Terbilang: {tb(x['jumlah'])} rupiah</i></p>{ttd('PPK')}</div>" for x in sp)
    elif jenis=="kuitansi":b=f"<h3 align=center>KUITANSI ({html_escape(s['jenis'])})</h3>{ident}<p>Telah terima sejumlah <b>{rp(tot)}</b><br><i>Terbilang: {tb(tot)} rupiah</i><br>Untuk pembayaran perjalanan dinas {html_escape(s['maksud'])}</p>{ttd('BENDAHARA')}"
    elif jenis=="nominatif":b="<h3 align=center>DAFTAR NOMINATIF</h3>"+ident+tbl([(x["urutan"],x["nama"],rp(x["uang_harian"]),rp(x["uang_inap"]),rp(x["tiket"]),rp(x["taksi"]+x["lain"]+x["rep"]),rp(x["jumlah"])) for x in sp],["No","Nama","Harian","Inap","Tiket","Lain+Rep","Jumlah"])+f"<p><b>Total: {rp(tot)}</b></p>"+ttd("PEMBUAT_DAFTAR")
    else:err("404","Jenis dokumen tidak ada",404)
    return f"<html><body style='font-family:serif;max-width:800px;margin:auto'>{kop}{b}<script>print()</script></body></html>"
@app.get("/api/v1/doc/{i}/{jenis}",response_class=HTMLResponse)
def doc_v1(i:int,jenis:str,u=Depends(user)):
    return doc(i,jenis,u)

PDF_JENIS={
    "st","st_lampiran","spd","rincian","kuitansi","nominatif",
    "sptjb","laporan","riil","lengkap"
}

def tanggal_pdf(value):
    if isinstance(value,dt.datetime):return value.date()
    if isinstance(value,dt.date):return value
    return dt.date.fromisoformat(value) if value else None

def siapkan_data_pdf(s,pegawai_rows):
    nilai_satker={r["k"]:r["v"] for r in q("select k,v from satker")}
    satker=dict(nilai_satker)
    satker["nama_satker_panjang"]=satker.get("nama_satker_panjang") or satker.get("nama_satker")
    satker["kota_kedudukan"]=satker.get("kota_kedudukan") or satker.get("kota")
    satker["dipa_raw"]=satker.get("dipa_raw") or satker.get("nomor_dipa")
    satker["nomor_bukti"]=s.get("nomor_bukti")
    satker["tahun_anggaran"]=s.get("tahun")
    satker["nomor_sptjb"]=satker.get("nomor_sptjb")
    roles={}
    for pejabat in q("select * from pejabat where aktif=1"):
        if pejabat["peran"] in roles:
            err("409",f"Ada lebih dari satu pejabat aktif dengan peran {pejabat['peran']}")
        roles[pejabat["peran"]]=dict(pejabat,jabatan_lengkap=pejabat.get("jabatan_lengkap") or pejabat.get("jabatan"))
    tanggal_st=tanggal_pdf(s.get("tanggal_surat_tugas"))
    tanggal_ttd=tanggal_pdf(s.get("tgl_tanda_tangan"))
    if not tanggal_st or not tanggal_ttd:
        err("422","Lengkapi tanggal surat tugas dan tanggal tanda tangan sebelum membuat PDF",422)
    st_pdf=SuratTugasHitung(
        nomor_surat_tugas=s.get("nomor") or "",tanggal_surat_tugas=tanggal_st,
        maksud_perjalanan=s.get("maksud") or "",asal=s.get("asal") or "",
        tujuan=s.get("tujuan") or "",provinsi_tarif=s.get("provinsi") or "",
        tgl_berangkat=tanggal_pdf(s["tgl_berangkat"]),
        tgl_kembali=tanggal_pdf(s["tgl_kembali"]),tgl_tanda_tangan=tanggal_ttd,
        akun_anggaran=s.get("akun") or "",nama_kegiatan=s.get("nama_kegiatan"),
        alat_angkutan=s.get("alat"),jenis_kuitansi=s.get("jenis"),status=s["status"],
        uraian_kegiatan=s.get("uraian_kegiatan"),keterangan_laporan=s.get("keterangan_laporan")
    )
    hasil=[]
    for row in pegawai_rows:
        if not all(row.get(k) for k in("nama","nip","golongan","jabatan")):
            err("422","Data nama, NIP, golongan, dan jabatan pelaksana harus lengkap sebelum mencetak",422)
        peg=PegawaiHitung(
            id=str(row.get("pegawai_id") or ""),nama=row.get("nama") or "",
            nip_nrp=row.get("nip") or "",golongan=row.get("golongan") or "",
            pangkat=row.get("pangkat") or "",jabatan=row.get("jabatan") or "",
            eselon=row.get("eselon"),tingkat_biaya=row.get("tingkat_biaya")
        )
        hasil.append(SPDHasil(
            pegawai=peg,hari=int(row["hari"]),tarif_uang_harian=int(row["tarif_harian"] or 0),
            uang_harian=int(row["uang_harian"] or 0),malam_inap=int(row["malam"] or 0),
            tarif_inap=int(row["tarif_inap"] or 0),uang_penginapan=int(row["uang_inap"] or 0),
            transport_tiket=int(row["tiket"] or 0),taksi=int(row["taksi"] or 0),
            biaya_lain=int(row["lain"] or 0),tarif_representasi=int(row["tarif_rep"] or 0),
            uang_representasi=int(row["rep"] or 0),jumlah=int(row["jumlah"] or 0),
            batas_sbm_inap=int(row["batas"] or 0),flag_lebih_sbm=bool(row["flag"]),
            ppn=row.get("ppn"),pph=row.get("pph")
        ))
    return satker,roles,st_pdf,hasil

def buat_pdf_dokumen(jenis,folder,s,satker,roles,st_pdf,hasil):
    paths=[]
    def path(nama):
        p=os.path.join(folder,f"{len(paths):03d}_{nama}.pdf")
        paths.append(p)
        return p
    peran={
        "st":["SEKDITJEN","PPK","KABAG"],
        "st_lampiran":["PPK","KABAG","SEKDITJEN"],
        "spd":["PPK","PENGESAH"],"rincian":["BENDAHARA","PPK"],
        "kuitansi":["BENDAHARA","PPK","PELAKSANA_ADM"],"nominatif":["PPK","PEMBUAT_DAFTAR","BENDAHARA"],
        "sptjb":["PPK"],"riil":["PPK"]
    }
    jenis_dihasilkan=[
        "st","st_lampiran","spd","rincian","kuitansi","nominatif",
        "sptjb","laporan","riil"
    ] if jenis=="lengkap" else [jenis]
    required_satker={
        "st":["nama_kementerian","nama_ditjen","nama_satker_panjang","alamat","kota_kedudukan"],
        "st_lampiran":["nama_satker"],
        "spd":["nama_kementerian","nama_ditjen","nama_satker_panjang","alamat","nama_satker","kota_kedudukan"],
        "rincian":["nama_satker","kota_kedudukan"],
        "kuitansi":["nama_kementerian","nama_ditjen","nama_satker_panjang","alamat","tahun_anggaran","kota_kedudukan"],
        "nominatif":["nama_satker","kota_kedudukan","rekening_bendahara"],
        "sptjb":["nama_satker","kota_kedudukan","nomor_sptjb","kode_satker","dipa_raw","klasifikasi_anggaran"],
        "laporan":["kota_kedudukan"],
        "riil":["nama_satker","kota_kedudukan"]
    }
    for tipe in jenis_dihasilkan:
        if any(not satker.get(k) for k in required_satker[tipe]):
            err("422",f"Identitas satker untuk dokumen {tipe} belum lengkap",422)
        if tipe=="st" and not q("select 1 from dasar_hukum where aktif=1"):
            err("422","Tambahkan dasar hukum aktif sebelum mencetak Surat Tugas",422)
    if "spd" in jenis_dihasilkan or "lengkap" in jenis_dihasilkan:
        if any(not h.pegawai.pangkat or not h.pegawai.tingkat_biaya for h in hasil):
            err("422","Lengkapi pangkat dan tingkat biaya perjalanan dinas pada data pegawai",422)
    for tipe in jenis_dihasilkan:
        for role in peran.get(tipe,[]):
            p=roles.get(role)
            if not p or not p.get("nama") or not p.get("nip") or set(str(p["nip"]))=={"0"} or not p.get("jabatan_lengkap"):
                err("422",f"Data pejabat '{role}' belum lengkap atau masih berupa data contoh",422)
        if tipe=="st":
            buat_pdf_st(path(tipe),satker,roles,q("select urutan,teks from dasar_hukum where aktif=1 order by urutan"),st_pdf)
        elif tipe=="st_lampiran":
            buat_pdf_st_lampiran(path(tipe),satker,roles,st_pdf,[h.pegawai for h in hasil])
        elif tipe in("spd","rincian","riil"):
            for h in hasil:
                if tipe=="spd":
                    f1=path(f"spd_{h.pegawai.id}_hal1")
                    buat_pdf_spd(f1,satker,roles["PPK"],st_pdf,h)
                    f2=path(f"spd_{h.pegawai.id}_hal2")
                    buat_pdf_spd_hal2(f2,satker,roles["PPK"],roles["PENGESAH"],st_pdf,h)
                elif tipe=="rincian":
                    buat_pdf_rincian(path(f"rincian_{h.pegawai.id}"),satker,roles["BENDAHARA"],roles["PPK"],st_pdf,h)
                else:
                    row=next(x for x in s["spd"] if str(x["pegawai_id"])==h.pegawai.id)
                    if row.get("biaya_riil") is None:
                        err("422",f"Nilai pengeluaran riil belum diisi untuk {h.pegawai.nama}",422)
                    buat_pdf_riil(path(f"riil_{h.pegawai.id}"),satker,roles["PPK"],st_pdf,h,int(row["biaya_riil"]))
        elif tipe=="kuitansi":
            if not satker.get("tahun_anggaran") or not satker.get("nomor_bukti"):
                err("422","Nomor bukti dan tahun anggaran wajib tersedia untuk kuitansi",422)
            buat_pdf_kuitansi(
                path(tipe),satker,roles["PPK"],roles["BENDAHARA"],
                roles["PELAKSANA_ADM"],st_pdf,
                sum(h.jumlah for h in hasil),satker["nomor_bukti"]
            )
        elif tipe=="nominatif":
            if not satker.get("rekening_bendahara"):err("422","Rekening bendahara belum diisi",422)
            buat_pdf_nominatif(path(tipe),satker,roles,st_pdf,hasil)
        elif tipe=="sptjb":
            if not satker.get("nomor_sptjb") or not satker.get("kode_satker") or not satker.get("dipa_raw") or not satker.get("klasifikasi_anggaran"):
                err("422","Data nomor SPTJB, kode satker, DIPA, dan klasifikasi anggaran wajib dilengkapi",422)
            if any(h.ppn is None or h.pph is None for h in hasil):
                err("422","Nilai PPN dan PPh aktual setiap pelaksana harus diisi sebelum menerbitkan SPTJB",422)
            buat_pdf_sptjb(path(tipe),satker,roles["PPK"],st_pdf,hasil)
        elif tipe=="laporan":
            buat_pdf_laporan(path(tipe),satker,st_pdf,[h.pegawai for h in hasil])
    if not paths:err("422","Tidak ada PDF yang dihasilkan")
    if len(paths)==1:
        with open(paths[0],"rb") as pdf:return pdf.read()
    output=os.path.join(folder,"gabungan.pdf")
    gabungkan_pdf(paths,output)
    with open(output,"rb") as pdf:return pdf.read()

@app.get("/api/st/{i}/pdf/{jenis}")
@app.get("/api/v1/st/{i}/pdf/{jenis}")
def pdf_st(i:int,jenis:str,u=Depends(user)):
    if jenis not in PDF_JENIS:err("404","Jenis PDF tidak tersedia",404)
    s=std(i,u)
    if s["status"] not in("TERBIT","SELESAI"):
        err("V8","Dokumen hanya bisa dicetak bila status TERBIT/SELESAI",403)
    if u["role"]=="PEGAWAI":
        if jenis not in("spd","rincian","riil"):err("403","Dokumen ini tidak tersedia untuk peran pegawai",403)
        s["spd"]=[x for x in s["spd"] if x["pegawai_id"]==u["pegawai_id"]]
        if not s["spd"]:err("403","Bukan SPD Anda",403)
    satker,roles,st_pdf,hasil=siapkan_data_pdf(s,s["spd"])
    with tempfile.TemporaryDirectory(prefix="pejadin-pdf-") as folder:
        content=buat_pdf_dokumen(jenis,folder,s,satker,roles,st_pdf,hasil)
    log(u,"CETAK","st",i,jenis)
    ex(
        "insert into dokumen_terbit(surat_tugas_id,spd_id,jenis_dokumen,dicetak_oleh) values(?,?,?,?)",
        (i,s["spd"][0]["id"] if jenis in("spd","rincian","riil") and len(s["spd"])==1 else None,jenis,u["id"])
    )
    return Response(
        content=content,media_type="application/pdf",
        headers={"Content-Disposition":f'attachment; filename="pejadin_{i}_{jenis}.pdf"'}
    )

@app.get("/api/laporan")
def lap(format:str="xlsx",tahun:int=0,status:str="",u=Depends(need("OPERATOR","PPK","BENDAHARA"))):
    rows=[r for r in stl(status,"",u) if not tahun or r["tahun"]==tahun];H=["No ST","Tujuan","Berangkat","Kembali","Pelaksana","Total Biaya","Status"]
    R=[(r["nomor"] or "-",r["tujuan"],r["tgl_berangkat"],r["tgl_kembali"],r["jml"],r["total"],r["status"]) for r in rows]
    if format=="pdf":
        display_rows=[(*row[:5],rp(row[5]),row[6]) for row in R]
        return HTMLResponse("<html><body style='font-family:sans-serif'><h3>Rekap Surat Tugas</h3><table border=1 cellpadding=4 style='border-collapse:collapse'><tr>"+"".join(f"<th>{h}</th>" for h in H)+"</tr>"+"".join("<tr>"+"".join(f"<td>{c}</td>" for c in row)+"</tr>" for row in display_rows)+"</table><script>print()</script></body></html>")
    from openpyxl import Workbook
    w=Workbook();a=w.active;a.append(H)
    for row in R:
        a.append(row)
        a.cell(a.max_row,6).number_format='"Rp" #,##0;-"Rp" #,##0;"Rp" 0'
    o=io.BytesIO();w.save(o);o.seek(0)
    return StreamingResponse(o,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=rekap_surat_tugas.xlsx"})
@app.get("/api/v1/laporan")
def lap_v1(format:str="xlsx",tahun:int=0,status:str="",u=Depends(need("OPERATOR","PPK","BENDAHARA"))):
    return lap(format,tahun,status,u)
FRONTEND=os.path.join(D,"..","frontend")
if os.path.isdir(FRONTEND):
    app.mount("/",StaticFiles(directory=FRONTEND,html=True))
